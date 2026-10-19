"""研究・高密度スライドのビルダー(python-pptx)。値はすべて tokens/*.json から読む。

    import sys; sys.path.insert(0, "lib")
    from research_deck import Deck
    deck = Deck(ratio="4:3")                       # "4:3" か "16:9"
    deck.cover("題", "所属・学会名", ["氏名", "2026-10-20"], ["要点1", "要点2", "要点3"])
    s = deck.slide("結論を1文で書いたタイトル", "何を比べたページか")
    x, w = s.col(0, 8)                              # 12 カラムのうち 0 番から 8 カラム分
    s.image("fig.png", x, s.top, w)
    x, w = s.col(8, 4)
    s.text(x, s.top, w, 2.0, ["読み取り {accent|強調} [1]", "2行目"])
    x, w = s.col(0, 12)
    s.table(x, 5.2, w, [["手法", "RMSE [eV]"], ["A", "0.12"]])
    x, w = s.col(0, 7)
    s.equation(r"y = f(\\mathbf{x}) + \\varepsilon", x, 3.0, w, 0.7, number=1)
    s.sources(["[1] Author, A. et al. J. Abbr. 2024, 1, 1."])
    deck.references(["Author, A. A.; ... https://doi.org/..."])
    deck.save("out.pptx")

文字の中の {accent|語} は accent、{alert|語} は alert、{bold|語} は W6 になる。
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))


def _load(name):
    return json.loads((ROOT / "tokens" / name).read_text(encoding="utf-8"))


COLORS = _load("colors.json")
TYPO = _load("typography.json")
TABLE = _load("tables.json")
FONT = TYPO["family"]["pptx"]["regular"]  # PowerPoint(Mac)はウェイト付きの名前でないとヒラギノを見つけられない
FONT_BOLD = TYPO["family"]["pptx"]["bold"]
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # 表スタイル「スタイルなし、表のグリッドなし」


def color(name: str) -> str:
    """トークン名(ink, accent …)か #RRGGBB を 6 桁の HEX にする。"""
    if name.startswith("#"):
        return name[1:]
    return COLORS[name]["value"][1:]


def _style(name):
    st = TYPO["styles"][name]
    return st["size"], st["weight"] >= 600, st["color"]


_MARK = re.compile(r"\{(accent|alert|bold)\|([^{}]*)\}")


def _set_font(run, size, bold, col, family=None):
    family = family or (FONT_BOLD if bold else FONT)
    run.font.size = Pt(size)
    run.font.bold = False  # 太さは W6 の書体で出す(擬似太字にしない)
    run.font.color.rgb = RGBColor.from_string(color(col))
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", family)


def _fill_paragraph(p, line: str, style: str, col: str | None, size_override=None):
    size, bold, base = _style(style)
    size = size_override or size
    base = col or base
    pos = 0
    for m in _MARK.finditer(line):
        if m.start() > pos:
            _set_font(p.add_run(), size, bold, base)
            p.runs[-1].text = line[pos:m.start()]
        kind, word = m.groups()
        r = p.add_run()
        r.text = word
        _set_font(r, size, bold or kind == "bold", kind if kind != "bold" else base)
        pos = m.end()
    if pos < len(line) or not p.runs:
        r = p.add_run()
        r.text = line[pos:]
        _set_font(r, size, bold, base)


def _no_margin(tf, lr=0.0, tb=0.0):
    tf.margin_left = tf.margin_right = Inches(lr)
    tf.margin_top = tf.margin_bottom = Inches(tb)


class Slide:
    def __init__(self, deck: "Deck", slide, number: int):
        self.deck, self.s, self.index = deck, slide, number
        sp = deck.spacing
        self.margin = sp["margin_in"]
        self.top = sp["zones"]["content"]["top_in"]
        self.bottom = sp["zones"]["content"]["bottom_in"]
        self.col_w = sp["grid"]["column_in"]
        self.gutter = sp["grid"]["gutter_in"]
        self.width = sp["slide"]["width_in"] - 2 * self.margin

    # ---- グリッド ----
    def col(self, start: int, span: int):
        """12 カラムのうち start 番(0 始まり)から span 本ぶんの (x, w) [in]。"""
        x = self.margin + start * (self.col_w + self.gutter)
        w = span * self.col_w + (span - 1) * self.gutter
        if start + span == 12:  # 端数は最終カラムで吸収
            w = self.margin + self.width - x
        return x, w

    # ---- 文字 ----
    def text(self, x, y, w, h, lines, style="body", col=None, align="left", anchor="top", spacing=None, size=None):
        """文字を置く。size で pt を上書きできる(余白があるときに大きくする。上限は typography.json の max)。"""
        if size:
            mx = TYPO["styles"][style].get("max", size)
            size = min(size, mx)
        tb = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        _no_margin(tf)
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
        lh = TYPO["styles"][style].get("line_height", 1.3)
        for i, line in enumerate([lines] if isinstance(lines, str) else lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
            p.line_spacing = Pt((size or TYPO["styles"][style]["size"]) * lh)  # 倍率指定はヒラギノで行間が広がりすぎるので pt で固定
            if spacing:
                p.space_before = Pt(spacing if i else 0)
            _fill_paragraph(p, line, style, col, size)
        return tb

    def bullets(self, x, y, w, h, items, style="body", size=None, spacing=6):
        return self.text(x, y, w, h, ["・" + t for t in items], style, spacing=spacing, size=size)

    def heading(self, x, y, w, text, size=None):
        return self.text(x, y, w, 0.34, text, "heading", size=size)

    def number(self, x, y, w, label, value, note="", accent=False, size=None, label_size=None):
        """要点数値: ラベル(label)→ 数値(number)→ 比較対象や条件(label)。戻り値は次に置ける y [in]。"""
        vs = size or TYPO["styles"]["number"]["size"]
        ls = label_size or TYPO["styles"]["label"]["size"]
        lh = ls * 1.35 / 72
        vh = vs * 1.2 / 72
        self.text(x, y, w, lh, label, "label", size=ls)
        self.text(x, y + lh, w, vh, value, "number", col="accent" if accent else None, size=vs)
        if note:
            self.text(x, y + lh + vh, w, lh, note, "label", size=ls)
        return y + lh + vh + (lh if note else 0)

    def rule(self, x, y, w, col="line-strong", pt=0.75):
        ln = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
        ln.line.color.rgb = RGBColor.from_string(color(col))
        ln.line.width = Pt(pt)
        return ln

    def arrow(self, x1, y1, x2, y2, col="ink", pt=1.5):
        ln = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        ln.line.color.rgb = RGBColor.from_string(color(col))
        ln.line.width = Pt(pt)
        tail = etree.SubElement(ln.line._get_or_add_ln(), qn("a:tailEnd"))
        tail.set("type", "triangle")
        return ln

    # ---- 図・表・式 ----
    def image(self, path, x, y, w=None, h=None, credit=None):
        pic = self.s.shapes.add_picture(str(path), Inches(x), Inches(y),
                                        Inches(w) if w else None, Inches(h) if h else None)
        if credit:
            bottom = Emu(pic.top + pic.height).inches
            self.text(x, bottom + 0.04, Emu(pic.width).inches, 0.24, credit, "label")
        return pic

    def table(self, x, y, w, rows, col_widths=None, align=None, row_h=None, size=None, bold_cells=()):
        """色なし・黒罫線の表(tokens/tables.json)。rows[0] が見出し行。

        align: 列ごとの "l" / "r" / "c"。省略時は 1 列目左、数値らしい列は右。
        bold_cells: 決め手のセル (行, 列) を W6 にする(色は使わない)。
        """
        nr, nc = len(rows), len(rows[0])
        row_h = row_h or TABLE["row_height_in"]
        gf = self.s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
        tbl = gf.table
        tblpr = tbl._tbl.tblPr
        for k in ("firstRow", "bandRow"):
            tblpr.set(k, "0")
        sid = tblpr.find(qn("a:tableStyleId"))
        if sid is None:
            sid = etree.SubElement(tblpr, qn("a:tableStyleId"))
        sid.text = NO_STYLE_TABLE
        col_widths = col_widths or [w / nc] * nc
        for j, cw in enumerate(col_widths):
            tbl.columns[j].width = Inches(cw)
        for i in range(nr):
            tbl.rows[i].height = Inches(row_h)
        if align is None:
            num = re.compile(r"^[-+−±0-9.,%\s()×]+$")
            align = ["l"] + ["r" if all(num.match(str(r[j]) or "0") for r in rows[1:]) else "l" for j in range(1, nc)]
        size_pt = size or TYPO["styles"]["table"]["size"]
        m = TABLE["cell_margin_in"]
        for i, r in enumerate(rows):
            for j, val in enumerate(r):
                c = tbl.cell(i, j)
                c.fill.background()
                c.margin_left = c.margin_right = Inches(m["left_right"])
                c.margin_top = c.margin_bottom = Inches(m["top_bottom"])
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = c.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.alignment = {"l": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT, "c": PP_ALIGN.CENTER}[align[j] if i else ("l" if align[j] == "l" else "c")]
                run = p.add_run()
                run.text = str(val)
                _set_font(run, size_pt, i == 0 or (i, j) in bold_cells, "ink")
                top = 1.0 if i <= 1 else 0.5  # 上端と見出し下は 1.0pt
                bot = 1.0 if i == nr - 1 or i == 0 else 0.5
                self._cell_border(c, "a:lnT", top)
                self._cell_border(c, "a:lnB", bot)
                self._cell_border(c, "a:lnL", 0.5 if j > 0 else None)
                self._cell_border(c, "a:lnR", 0.5 if j < nc - 1 else None)
        return gf

    @staticmethod
    def _cell_border(cell, tag, pt):
        tcpr = cell._tc.get_or_add_tcPr()
        old = tcpr.find(qn(tag))
        if old is not None:
            tcpr.remove(old)
        ln = etree.Element(qn(tag))
        if pt is None:
            ln.set("w", "0")
            etree.SubElement(ln, qn("a:noFill"))
        else:
            ln.set("w", str(int(pt * 12700)))
            ln.set("cap", "flat")
            ln.set("cmpd", "sng")
            fill = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(fill, qn("a:srgbClr")).set("val", color("ink"))
            etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")
        # 罫線は塗り(fill)より前に置く決まり
        order = ["a:lnL", "a:lnR", "a:lnT", "a:lnB"]
        idx = 0
        for k, child in enumerate(tcpr):
            if child.tag in [qn(o) for o in order[: order.index(tag)]]:
                idx = k + 1
        tcpr.insert(idx, ln)

    def equation(self, latex, x, y, w, h, number=None, size=None):
        """別行立ての式。Word の XSL があればネイティブ数式、なければ $$…$$ のプレースホルダー。

        どちらも図形の代替テキストに元の LaTeX を残す(tools/latex_eq.py --extract で取り出せる)。
        """
        from latex_eq import place_equation
        size = size or TYPO["styles"]["equation"]["size"]
        return place_equation(self.s, latex, x, y, w, h, size, number)

    def notes(self, text: str):
        """発表者ノート(話す内容・補足・LaTeX の原文など)。"""
        self.s.notes_slide.notes_text_frame.text = text

    # ---- 出典欄 ----
    def sources(self, items):
        """出典欄(2行まで)。そのページで引いた文献の短縮形を続けて書く。"""
        z = self.deck.spacing["zones"]["footer"]
        w = self.width - 0.8
        self.text(self.margin, z["top_in"], w, z["bottom_in"] - z["top_in"], "  ".join(items), "caption")


class Deck:
    def __init__(self, ratio: str = "4:3"):
        self.ratio = ratio
        self.spacing = _load("spacing.json" if ratio == "4:3" else "spacing-16x9.json")
        self.prs = Presentation()
        self.prs.slide_width = Inches(self.spacing["slide"]["width_in"])
        self.prs.slide_height = Inches(self.spacing["slide"]["height_in"])
        self._layout = self.prs.slide_layouts[6]
        self.count = 0

    def blank(self) -> Slide:
        """タイトル帯のないページ(表紙など自由に組むとき)。ページ番号だけ入る。"""
        return self._new()

    def _new(self) -> Slide:
        self.count += 1
        s = Slide(self, self.prs.slides.add_slide(self._layout), self.count)
        z = self.spacing["zones"]["footer"]
        s.text(s.margin + s.width - 0.8, z["top_in"], 0.8, 0.2, str(self.count), "caption", align="right")
        return s

    def slide(self, title: str, message: str = "") -> Slide:
        s = self._new()
        z = self.spacing["zones"]
        s.text(s.margin, z["title"]["top_in"], s.width, z["title"]["height_in"] - 0.15, title, "title", anchor="top")
        if message:
            s.text(s.margin, z["message"]["top_in"], s.width, z["message"]["height_in"], message, "message")
        s.rule(s.margin, z["rule"]["y_in"], s.width, "line-strong", z["rule"]["height_pt"])
        return s

    def cover(self, title, subtitle, author_lines, points, sources=None) -> Slide:
        s = self._new()
        s.rule(s.margin, 2.6, 0.8, "ink", 0.75)
        s.text(s.margin, 2.72, s.width, 1.3, title, "cover-title")
        s.text(s.margin, 4.05, s.width, 0.8, subtitle, "cover-sub")
        left_span = 5 if self.ratio == "4:3" else 6
        x, w = s.col(0, left_span)
        s.text(x, 5.4, w, 1.4, author_lines, "body")
        x, w = s.col(left_span, 12 - left_span)
        s.heading(x, 5.4, w, "本発表の要点")
        s.text(x, 5.72, w, 1.13, ["・" + p for p in points], "body")
        if sources:
            s.sources(sources)
        return s

    def references(self, refs, title="参考文献", columns=None, extra_note=None) -> Slide:
        s = self.slide(title, "本文の登場順。DOI と Web 資料の確認日つき")
        columns = columns or (2 if self.ratio == "4:3" else 3)
        per = -(-len(refs) // columns)
        span = 12 // columns
        for c in range(columns):
            x, w = s.col(c * span, span)
            chunk = refs[c * per:(c + 1) * per]
            s.text(x, s.top, w, s.bottom - s.top, [f"[{c * per + i + 1}] {r}" for i, r in enumerate(chunk)], "caption", spacing=4)
        if extra_note:
            s.sources([extra_note])
        return s

    def save(self, path, title: str = ""):
        cp = self.prs.core_properties  # python-pptx の既定の作成者名を残さない
        cp.author = cp.last_modified_by = ""
        cp.title = title
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(path))
        return path


def use_chart_style():
    """matplotlib に tokens/charts.json の様式を当てる。系列の色・マーカー・線種のリストを返す。"""
    import matplotlib as mpl
    mpl.style.use(str(ROOT / "lib" / "research.mplstyle"))
    s = COLORS["series"]
    return list(zip(s["value"], s["markers"], s["linestyles"], COLORS["series-text"]["value"]))
