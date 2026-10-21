"""LaTeX の式を PowerPoint のネイティブ数式(OMML)としてスライドに挿入する。

経路: LaTeX → MathML(latex2mathml)→ OMML(Word 同梱の mathml2omml.xsl)→ a14:m としてテキストボックスへ。
挿入した式は PowerPoint の数式エディタでそのまま編集できる。画像では代用しない。

依存(環境を汚さず一時導入):
    uv run --no-project --with latex2mathml --with lxml --with python-pptx python latex_eq.py ...

CLI:
    latex_eq.py deck.pptx --slide 3 --latex "\\sigma_{LT} = c\\,\\sigma_{HT}^{n}" --x 1 --y 2 --w 6 --h 0.8 \
        [--size 18] [--number 1] [--align center|left] [--out out.pptx]
    latex_eq.py --check "\\frac{a}{b}"      # 変換できるかだけ確かめる(OMML を表示)
    latex_eq.py deck.pptx --convert          # LaTeX プレースホルダー($$ … $$)をまとめてネイティブ数式へ
    latex_eq.py deck.pptx --extract          # 全数式の LaTeX を一覧(代替テキストに保存してある)

どの式も図形の代替テキストに「LaTeX: …」として元の LaTeX を残す。Word の XSL がない環境では
place_equation() が $$ … $$ のプレースホルダーを置くので、後で Word のある環境で --convert する。

関数:
    from latex_eq import add_equation
    add_equation(slide, r"\\Lambda = \\Sigma^{-1}", x=1, y=2, w=6, h=0.8, size=18, number=1)
"""
from __future__ import annotations

import argparse
import copy
import sys
from pathlib import Path

from lxml import etree

XSL_CANDIDATES = [
    "/Applications/Microsoft Word.app/Contents/Resources/mathml2omml.xsl",
    "/Applications/Microsoft PowerPoint.app/Contents/Resources/mathml2omml.xsl",
    r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL",
    r"C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL",
]
NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "a14": "http://schemas.microsoft.com/office/drawing/2010/main",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
}
INK = "1A1A1A"
EQNUM = "404040"
EQNUM_FONT = "Hiragino Kaku Gothic ProN W3"
PLACEHOLDER_NAME = "latex-eq"  # 後で変換する LaTeX のプレースホルダー(図形名)
_XSLT = None


class LatexEquationError(ValueError):
    """変換できない LaTeX。画像に逃げず、ここで止める。"""


def _xslt():
    global _XSLT
    if _XSLT is None:
        path = next((p for p in XSL_CANDIDATES if Path(p).exists()), None)
        if path is None:
            raise FileNotFoundError("mathml2omml.xsl が見つからない(Microsoft Word が必要)。候補: " + ", ".join(XSL_CANDIDATES))
        _XSLT = etree.XSLT(etree.parse(path))
    return _XSLT


MML = "http://www.w3.org/1998/Math/MathML"
_SPACES = [(0.15, "\u200a"), (0.2, "\u2009"), (0.25, "\u205f"), (0.35, "\u2004"), (0.75, "\u2002"), (1.5, "\u2003"), (99, "\u2003\u2003")]


def _plain_alnum(ch: str):
    """数学用英数字(𝐲 など U+1D400–1D7FF)を (普通の文字, mathvariant) に戻す。PowerPoint はこれらを描けないため。"""
    cp = ord(ch)
    styles = [("bold", 0x1D400), ("italic", 0x1D434), ("bold-italic", 0x1D468), ("script", 0x1D49C), ("bold-script", 0x1D4D0),
              ("fraktur", 0x1D504), ("double-struck", 0x1D538), ("bold-fraktur", 0x1D56C), ("sans-serif", 0x1D5A0),
              ("sans-serif-bold", 0x1D5D4), ("sans-serif-italic", 0x1D608), ("sans-serif-bold-italic", 0x1D63C), ("monospace", 0x1D670)]
    for name, base in styles:
        if base <= cp < base + 52:
            off = cp - base
            return (chr(ord("A") + off) if off < 26 else chr(ord("a") + off - 26)), name
    if 0x1D7CE <= cp <= 0x1D7FF:  # 太字などの数字
        return chr(ord("0") + (cp - 0x1D7CE) % 10), "bold"
    return None


def _prep_mathml(root):
    """Word の mathml2omml.xsl が崩す形を、変換前に直す。"""
    q = lambda t: f"{{{MML}}}{t}"
    # 1) \left( ... \right) の中身を1つの mrow に包む(中の演算子が区切り文字 sepChr と誤認されるのを防ぐ)
    for row in list(root.iter(q("mrow"))):
        kids = list(row)
        if len(kids) >= 3 and kids[0].tag == q("mo") and kids[-1].tag == q("mo") \
                and kids[0].get("fence") == "true" and kids[-1].get("fence") == "true" and len(kids) > 3:
            inner = etree.Element(q("mrow"))
            for k in kids[1:-1]:
                inner.append(k)
            kids[0].addnext(inner)
    # 2) 行列の括弧: ( mtable ) を伸びる括弧にする
    for row in list(root.iter(q("mrow"))):
        kids = list(row)
        for i in range(len(kids) - 2):
            a, t, b = kids[i:i + 3]
            if a.tag == q("mo") and b.tag == q("mo") and t.tag == q("mtable"):
                for m, form in ((a, "prefix"), (b, "postfix")):
                    m.set("fence", "true"); m.set("stretchy", "true"); m.set("form", form)
                grp = etree.Element(q("mrow"))
                a.addprevious(grp)
                for m in (a, t, b):
                    grp.append(m)
    # 2b) 別行立てでは総和・総乗の添字を上下に置く(msubsup → munderover)
    if root.get("display") == "block":
        big = {"\u2211", "\u220f", "\u2210", "\u22c3", "\u22c2"}
        for tag, new in (("msubsup", "munderover"), ("msub", "munder"), ("msup", "mover")):
            for el in list(root.iter(q(tag))):
                base = el[0] if len(el) else None
                if base is not None and base.tag == q("mo") and (base.text or "").strip() in big:
                    el.tag = q(new)
    # 3) 数学用英数字を普通の文字 + mathvariant に戻す
    for el in root.iter(q("mi"), q("mn"), q("mtext")):
        txt = el.text or ""
        conv = [_plain_alnum(c) for c in txt]
        if txt and all(conv):
            el.text = "".join(c for c, _ in conv)
            el.set("mathvariant", conv[0][1])
    # 4) mspace(\, \quad など)を空白文字に置き換える(XSL は mspace を捨てる)
    for sp in list(root.iter(q("mspace"))):
        w = sp.get("width", "0.167em")
        try:
            em = float(w.replace("em", "")) if w.endswith("em") else 0.167
        except ValueError:
            em = 0.167
        ch = next(c for lim, c in _SPACES if em <= lim)
        mo = etree.Element(q("mtext")); mo.text = ch
        sp.addprevious(mo); sp.getparent().remove(sp)
    return root


def check_fraction_style(latex: str):
    """分数は横線(\\frac)で書く決まり。\\text{…} の外にスラッシュがあれば止める。"""
    import re
    body = re.sub(r"\\(text|mathrm|operatorname)\{[^{}]*\}", "", latex)
    if "/" in body:
        raise LatexEquationError(f"分数はスラッシュ(a/b)ではなく \\frac{{a}}{{b}}(横線)で書く: {latex!r}")


def _force_bar_fractions(math):
    """OMML の分数を横線(bar)にそろえる。斜線(skw)・一行(lin)の指定を消す。"""
    for t in math.iter(f"{{{NS['m']}}}type"):
        if t.get(f"{{{NS['m']}}}val") in ("skw", "lin"):
            t.getparent().remove(t)


def latex_to_omml(latex: str, display: bool = True) -> etree._Element:
    """LaTeX を m:oMath 要素にする(display=True で総和の添字を上下に置く別行立て)。失敗したら LatexEquationError。"""
    check_fraction_style(latex)
    try:
        import latex2mathml.converter as l2m
    except ImportError as e:  # pragma: no cover
        raise ImportError("latex2mathml が必要: uv run --no-project --with latex2mathml ...") from e
    depth = 0
    for ch in latex.replace(r"\{", "").replace(r"\}", ""):
        depth += (ch == "{") - (ch == "}")
        if depth < 0:
            break
    if depth != 0:
        raise LatexEquationError(f"波括弧 {{ }} の対応が取れていない: {latex!r}")
    try:
        mathml = l2m.convert(latex, display="block" if display else "inline")
    except Exception as e:
        raise LatexEquationError(f"LaTeX → MathML に失敗: {latex!r}: {e}") from e
    leftover = [t for t in etree.fromstring(mathml.encode()).itertext() if "\\" in t]
    if leftover:  # latex2mathml が解釈できなかった命令はバックスラッシュ付きの文字として残る
        raise LatexEquationError(f"未対応の命令 {leftover} を含む: {latex!r}")
    if "<merror" in mathml:
        raise LatexEquationError(f"未対応の記法を含む: {latex!r}")
    omml = _xslt()(_prep_mathml(etree.fromstring(mathml.encode())))
    root = omml.getroot()
    if root is None:
        raise LatexEquationError(f"MathML → OMML に失敗: {latex!r}")
    math = root if root.tag == f"{{{NS['m']}}}oMath" else root.find(".//m:oMath", NS)
    if math is None:
        raise LatexEquationError(f"OMML に m:oMath がない: {latex!r}")
    for sep in math.iter(f"{{{NS['m']}}}sepChr"):
        if not sep.get(f"{{{NS['m']}}}val"):
            sep.getparent().remove(sep)
    _force_bar_fractions(math)
    return math


def _run_props(size_pt: float, color: str, italic: bool | None):
    rpr = etree.Element(f"{{{NS['a']}}}rPr", lang="en-US", sz=str(int(round(size_pt * 100))))
    if italic is not None:
        rpr.set("i", "1" if italic else "0")
    fill = etree.SubElement(rpr, f"{{{NS['a']}}}solidFill")
    etree.SubElement(fill, f"{{{NS['a']}}}srgbClr", val=color)
    for tag in ("latin", "ea", "cs"):
        etree.SubElement(rpr, f"{{{NS['a']}}}{tag}", typeface="Cambria Math")
    return rpr


def _to_drawingml_runs(math: etree._Element, size_pt: float, color: str):
    """Word 向けの w:rPr を外し、各 m:r に DrawingML の a:rPr(Cambria Math・サイズ・色)を入れる。"""
    for w in math.xpath(".//w:*", namespaces=NS):
        w.getparent().remove(w)
    for r in math.iter(f"{{{NS['m']}}}r"):
        mrpr = r.find("m:rPr", NS)
        sty = mrpr.find("m:sty", NS) if mrpr is not None else None
        sty = sty.get(f"{{{NS['m']}}}val") if sty is not None else None
        rpr = _run_props(size_pt, color, False if sty in ("p", "b") else None)
        if sty in ("b", "bi"):
            rpr.set("b", "1")
        t = r.find("m:t", NS)
        (t.addprevious(rpr) if t is not None else r.append(rpr))
    # 演算子・括弧などのプロパティにも書体を入れる(ctrlPr)
    for ctrl in math.iter(f"{{{NS['m']}}}ctrlPr"):
        for c in list(ctrl):
            ctrl.remove(c)
        ctrl.append(_run_props(size_pt, color, True))
    return math


def _math_paragraph(latex: str, size_pt: float, color: str, align: str):
    math = _to_drawingml_runs(latex_to_omml(latex), size_pt, color)
    para = etree.Element(f"{{{NS['m']}}}oMathPara", nsmap={"m": NS["m"]})
    ppr = etree.SubElement(para, f"{{{NS['m']}}}oMathParaPr")
    etree.SubElement(ppr, f"{{{NS['m']}}}jc", {f"{{{NS['m']}}}val": "centerGroup" if align == "center" else "left"})
    para.append(math)
    a14m = etree.Element(f"{{{NS['a14']}}}m", nsmap={"a14": NS["a14"]})
    a14m.append(para)
    return a14m


def add_equation(slide, latex: str, x: float, y: float, w: float, h: float, size: float = 18,
                 number: int | str | None = None, align: str = "center", color: str = INK):
    """スライドに別行立ての式を置く(単位 in)。number を渡すと右端に式番号 (n) を 12pt で付ける。

    戻り値は式のテキストボックス(python-pptx の Shape)。PowerPoint で開くと数式として編集できる。
    """
    from pptx.util import Inches, Pt

    num_w = 0.6 if number is not None else 0.0
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w - num_w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    from pptx.enum.text import MSO_ANCHOR
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    p = tf.paragraphs[0]._p
    for child in list(p):
        if child.tag != f"{{{NS['a']}}}pPr":
            p.remove(child)
    p.append(_math_paragraph(latex, size, color, align))
    _set_alt(tb, latex)
    # PowerPoint 2010 以降は a14 を解釈する。それ以前のアプリには LaTeX の文字列を代わりに見せる
    sp = tb._element
    fallback = copy.deepcopy(sp)
    fp = fallback.find(".//a:p", NS)
    for child in list(fp):
        if child.tag != f"{{{NS['a']}}}pPr":
            fp.remove(child)
    r = etree.SubElement(fp, f"{{{NS['a']}}}r")
    r.append(_run_props(size * 0.7, color, False))
    etree.SubElement(r, f"{{{NS['a']}}}t").text = latex
    ac = etree.Element(f"{{{NS['mc']}}}AlternateContent", nsmap={"mc": NS["mc"]})
    choice = etree.SubElement(ac, f"{{{NS['mc']}}}Choice", nsmap={"a14": NS["a14"]}, Requires="a14")
    fb = etree.SubElement(ac, f"{{{NS['mc']}}}Fallback")
    sp.addprevious(ac)
    choice.append(sp)
    fb.append(fallback)
    if number is not None:
        nb = slide.shapes.add_textbox(Inches(x + w - num_w), Inches(y), Inches(num_w), Inches(h))
        ntf = nb.text_frame
        ntf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(ntf, m, 0)
        from pptx.enum.text import PP_ALIGN
        para = ntf.paragraphs[0]
        para.alignment = PP_ALIGN.RIGHT
        run = para.add_run()
        run.text = f"({number})"
        run.font.size = Pt(12)
        from pptx.dml.color import RGBColor
        run.font.color.rgb = RGBColor.from_string(EQNUM)
        run.font.name = EQNUM_FONT
    return tb


def _set_alt(shape, latex: str):
    """図形の代替テキスト(descr)に元の LaTeX を残す。後から再変換・検索・論文への転記に使う。"""
    cnv = shape._element.find(".//p:cNvPr", NS)
    if cnv is not None:
        cnv.set("descr", "LaTeX: " + latex)


def xsl_available() -> bool:
    return any(Path(p).exists() for p in XSL_CANDIDATES)


def add_latex_placeholder(slide, latex: str, x: float, y: float, w: float, h: float, size: float = 18,
                          number: int | str | None = None):
    """変換環境(Word の XSL)がないときに、LaTeX をそのまま置く。図形名 latex-eq と代替テキストに LaTeX を保存し、
    `latex_eq.py --convert deck.pptx` でまとめてネイティブ数式に置き換えられる。"""
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    num_w = 0.6 if number is not None else 0.0
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w - num_w), Inches(h))
    tb.name = PLACEHOLDER_NAME
    _set_alt(tb, latex)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = "$$ " + latex + " $$"
    run.font.size = Pt(max(12, size * 0.7))
    run.font.name = "Menlo"
    run.font.color.rgb = RGBColor.from_string(INK)
    if number is not None:
        nb = slide.shapes.add_textbox(Inches(x + w - num_w), Inches(y), Inches(num_w), Inches(h))
        nb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        np_ = nb.text_frame.paragraphs[0]
        np_.alignment = PP_ALIGN.RIGHT
        r = np_.add_run()
        r.text = f"({number})"
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor.from_string(EQNUM)
        r.font.name = EQNUM_FONT
    return tb


def place_equation(slide, latex: str, x: float, y: float, w: float, h: float, size: float = 18,
                   number: int | str | None = None, align: str = "center"):
    """変換できればネイティブ数式、変換環境がなければ LaTeX プレースホルダーを置く。記法の誤りはエラーで止める。"""
    latex_to_mathml_checked(latex)
    if xsl_available():
        return add_equation(slide, latex, x, y, w, h, size, number, align)
    return add_latex_placeholder(slide, latex, x, y, w, h, size, number)


def latex_to_mathml_checked(latex: str):
    """LaTeX の記法だけを確かめる(Word がなくても動く)。"""
    from latex2mathml.converter import convert
    if latex.count("{") != latex.count("}"):
        raise LatexEquationError("波括弧の数が合わない: " + latex)
    check_fraction_style(latex)
    try:
        return convert(latex)
    except Exception as e:  # noqa: BLE001
        raise LatexEquationError(f"LaTeX を解釈できない: {latex} ({e})") from e


def convert_placeholders(pptx_path, out=None) -> int:
    """プレースホルダー(latex-eq)をネイティブ数式に置き換える。式番号の図形はそのまま残す。"""
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation(str(pptx_path))
    n = 0
    for slide in prs.slides:
        for shp in list(slide.shapes):
            if shp.name != PLACEHOLDER_NAME:
                continue
            latex = shp.text_frame.text.strip()
            if latex.startswith("$$") and latex.endswith("$$"):
                latex = latex[2:-2].strip()
            size = shp.text_frame.paragraphs[0].runs[0].font.size.pt / 0.7 if shp.text_frame.paragraphs[0].runs else 18
            x, y, w, h = (Emu(v).inches for v in (shp.left, shp.top, shp.width, shp.height))
            add_equation(slide, latex, x, y, w, h, round(size), None, "center")
            shp._element.getparent().remove(shp._element)
            n += 1
    prs.save(str(out or pptx_path))
    return n


def extract_latex(pptx_path):
    """デッキ内の全数式の LaTeX を (スライド番号, LaTeX) で返す(論文・Markdown への転記用)。"""
    from pptx import Presentation
    found = []
    for i, slide in enumerate(Presentation(str(pptx_path)).slides, 1):
        for el in slide._element.iter(f"{{{NS['p']}}}cNvPr"):
            d = el.get("descr", "")
            if d.startswith("LaTeX: "):
                found.append((i, d[len("LaTeX: "):]))
    # AlternateContent の Choice と Fallback で同じ式が 2 回出るので重複を除く
    seen, out = set(), []
    for item in found:
        if item not in seen:
            seen.add(item)
            out.append(item)
    return out


def count_equations(pptx_path) -> int:
    """pptx 内のネイティブ数式(a14:m)の数を数える(検証用)。"""
    from pptx import Presentation
    n = 0
    for s in Presentation(str(pptx_path)).slides:
        n += sum(1 for _ in s._element.iter(f"{{{NS['a14']}}}m"))
    return n


def main(argv=None):
    ap = argparse.ArgumentParser(description="LaTeX → PowerPoint ネイティブ数式")
    ap.add_argument("pptx", nargs="?")
    ap.add_argument("--slide", type=int, help="1 始まりのスライド番号")
    ap.add_argument("--latex")
    ap.add_argument("--x", type=float, default=1.0)
    ap.add_argument("--y", type=float, default=2.0)
    ap.add_argument("--w", type=float, default=8.0)
    ap.add_argument("--h", type=float, default=0.8)
    ap.add_argument("--size", type=float, default=18)
    ap.add_argument("--number")
    ap.add_argument("--align", choices=["center", "left"], default="center")
    ap.add_argument("--out")
    ap.add_argument("--check", metavar="LATEX", help="変換だけ試して OMML を表示する")
    ap.add_argument("--convert", action="store_true", help="pptx 内の LaTeX プレースホルダーをネイティブ数式に置き換える")
    ap.add_argument("--extract", action="store_true", help="pptx 内の全数式の LaTeX を一覧表示する")
    a = ap.parse_args(argv)
    try:
        if a.check:
            print(etree.tostring(latex_to_omml(a.check), pretty_print=True).decode())
            return 0
        if a.convert and a.pptx:
            n = convert_placeholders(a.pptx, a.out)
            print(f"{a.out or a.pptx}: プレースホルダー {n} 式をネイティブ数式に変換")
            return 0
        if a.extract and a.pptx:
            for i, tex in extract_latex(a.pptx):
                print(f"slide {i}: $$ {tex} $$")
            return 0
        if not (a.pptx and a.slide and a.latex):
            ap.error("pptx, --slide, --latex が必要")
        from pptx import Presentation
        prs = Presentation(a.pptx)
        add_equation(prs.slides[a.slide - 1], a.latex, a.x, a.y, a.w, a.h, a.size, a.number, a.align)
        out = a.out or a.pptx
        prs.save(out)
        print(f"{out}: slide {a.slide} に数式を追加({count_equations(out)} 式)")
        return 0
    except LatexEquationError as e:
        print(f"エラー: {e}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
